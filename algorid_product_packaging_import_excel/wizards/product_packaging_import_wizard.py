# -*- encoding: utf-8 -*-
#################################################################
#                                                               #
#       Developed by Algorid Limited                            #
#       Copyright (C) 2025 - Today (https://www.algorid.com)    #
#       Contact us: contact@algorid.com                         #
#                                                               #
#################################################################

from odoo import models, fields, api
from odoo.exceptions import UserError, ValidationError
import base64
import io
import xlrd
from datetime import datetime


class ProductPackagingImportWizard(models.TransientModel):
    _name = 'product.packaging.import.wizard'
    _description = 'Import Product Packaging from Excel'

    file = fields.Binary(
        string='Excel File',
        help="Upload an Excel file (.xls or .xlsx) with packaging data"
    )
    filename = fields.Char(string='Filename')
    replace_existing = fields.Boolean(
        string='Replace Existing Packaging',
        default=False,
        help="If checked, existing packaging on products will be deleted before import"
    )
    company_id = fields.Many2one(
        'res.company',
        string='Company',
        default=lambda self: self.env.company
    )
    import_type = fields.Selection([
        ('create', 'Create New Packaging'),
        ('update', 'Update Existing Packaging'),
        ('both', 'Create or Update')
    ], string='Import Type', default='both', required=True)
    
    # Fields for product filtering
    filter_product_id = fields.Many2one(
        'product.product',
        string='Filter by Product',
        help="If set, only packaging for this product will be imported"
    )
    filter_product_tmpl_id = fields.Many2one(
        'product.template',
        string='Filter by Product Template',
        help="If set, only packaging for products in this template will be imported"
    )
    show_filter_info = fields.Boolean(
        compute='_compute_show_filter_info'
    )
    filter_info_text = fields.Char(
        compute='_compute_show_filter_info'
    )

    @api.depends('filter_product_id', 'filter_product_tmpl_id')
    def _compute_show_filter_info(self):
        """Compute filter information display"""
        for record in self:
            if record.filter_product_id:
                record.show_filter_info = True
                record.filter_info_text = f"Importing only for: {record.filter_product_id.display_name}"
            elif record.filter_product_tmpl_id:
                record.show_filter_info = True
                variant_count = len(record.filter_product_tmpl_id.product_variant_ids)
                record.filter_info_text = f"Importing only for: {record.filter_product_tmpl_id.display_name} ({variant_count} variant(s))"
            else:
                record.show_filter_info = False
                record.filter_info_text = False

    def action_download_template(self):
        """Download Excel template"""
        self.ensure_one()
        
        # Create template using openpyxl
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            raise UserError("Please install openpyxl: pip install openpyxl")
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Packaging Import"
        
        # Header style
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        # Define headers
        headers = [
            'Product Code/ID*',
            'Packaging Name*',
            'Sequence',
            'Contained Quantity*',
            'Barcode',
            'Package Type',
            'Sales',
            'Purchase',
            'Routes (comma separated)'
        ]
        
        # Write headers
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
        
        # Add example data based on filter
        example_data = []
        
        if self.filter_product_id:
            # Add examples for the specific product
            product = self.filter_product_id
            product_ref = product.default_code or f"[{product.id}]"
            example_data = [
                [product_ref, 'Box of 10', '1', '10', '1234567890', 'Box', 'TRUE', 'TRUE', 'Buy'],
                [product_ref, 'Pallet of 100', '2', '100', '0987654321', 'Pallet', 'TRUE', 'FALSE', 'Buy,Manufacture'],
            ]
        elif self.filter_product_tmpl_id:
            # Add examples for each variant
            for idx, variant in enumerate(self.filter_product_tmpl_id.product_variant_ids[:3], 1):
                product_ref = variant.default_code or f"[{variant.id}]"
                example_data.append([product_ref, f'Box of 10', str(idx), '10', '', 'Box', 'TRUE', 'TRUE', ''])
        else:
            # Default examples
            example_data = [
                ['PROD001', 'Box of 10', '1', '10', '1234567890', 'Box', 'TRUE', 'TRUE', 'Buy'],
                ['PROD001', 'Pallet of 100', '2', '100', '0987654321', 'Pallet', 'TRUE', 'FALSE', 'Buy,Manufacture'],
                ['[123]', 'Case of 6', '1', '6', '', 'Case', 'TRUE', 'TRUE', ''],
            ]
        
        for row_num, row_data in enumerate(example_data, 2):
            for col_num, value in enumerate(row_data, 1):
                ws.cell(row=row_num, column=col_num, value=value)
        
        # Add instructions sheet
        ws_info = wb.create_sheet("Instructions")
        instructions = [
            ["Field", "Description", "Required", "Format/Options"],
            ["Product Code/ID", "Product default code or ID in brackets [123]", "Yes", "Text or [ID]"],
            ["Packaging Name", "Name of the packaging", "Yes", "Text"],
            ["Sequence", "Order of packaging (lower = default)", "No", "Number (default: 1)"],
            ["Contained Quantity", "Quantity in this packaging", "Yes", "Number"],
            ["Barcode", "Barcode for the packaging", "No", "Text"],
            ["Package Type", "Package type name", "No", "Text (must exist)"],
            ["Sales", "Available for sales", "No", "TRUE/FALSE (default: TRUE)"],
            ["Purchase", "Available for purchase", "No", "TRUE/FALSE (default: TRUE)"],
            ["Routes", "Route names separated by comma", "No", "Text (comma separated)"],
            ["", "", "", ""],
            ["Notes:", "", "", ""],
            ["- Fields marked with * are required", "", "", ""],
            ["- Product Code/ID: Use product default_code or database ID in brackets like [123]", "", "", ""],
            ["- Boolean fields accept: TRUE/FALSE, YES/NO, 1/0, T/F, Y/N (case insensitive)", "", "", ""],
            ["- Leave Sequence empty to use default value (1)", "", "", ""],
            ["- Multiple routes should be separated by comma: Buy,Manufacture", "", "", ""],
        ]
        
        # Add filter information to instructions
        if self.show_filter_info:
            instructions.insert(11, ["", "", "", ""])
            instructions.insert(12, [f"FILTER ACTIVE: {self.filter_info_text}", "", "", ""])
            instructions.insert(13, ["Only packaging for the filtered product(s) will be imported!", "", "", ""])
        
        for row_num, row_data in enumerate(instructions, 1):
            for col_num, value in enumerate(row_data, 1):
                cell = ws_info.cell(row=row_num, column=col_num, value=value)
                if row_num == 1:
                    cell.fill = header_fill
                    cell.font = header_font
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 20
        ws.column_dimensions['E'].width = 15
        ws.column_dimensions['F'].width = 15
        ws.column_dimensions['G'].width = 10
        ws.column_dimensions['H'].width = 12
        ws.column_dimensions['I'].width = 30
        
        ws_info.column_dimensions['A'].width = 20
        ws_info.column_dimensions['B'].width = 50
        ws_info.column_dimensions['C'].width = 12
        ws_info.column_dimensions['D'].width = 40
        
        # Save to bytes
        excel_file = io.BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)
        
        # Generate filename
        filename = 'product_packaging_import_template'
        if self.filter_product_id:
            safe_name = self.filter_product_id.name.replace(' ', '_').replace('/', '_')[:30]
            filename = f'packaging_import_{safe_name}'
        elif self.filter_product_tmpl_id:
            safe_name = self.filter_product_tmpl_id.name.replace(' ', '_').replace('/', '_')[:30]
            filename = f'packaging_import_{safe_name}'
        
        # Return download action
        attachment = self.env['ir.attachment'].create({
            'name': f'{filename}.xlsx',
            'type': 'binary',
            'datas': base64.b64encode(excel_file.read()),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }

    def action_import_packaging(self):
        """Import packaging from Excel file"""
        self.ensure_one()
        
        if not self.file:
            raise UserError("Please upload an Excel file.")
        
        # Decode file
        try:
            file_content = base64.b64decode(self.file)
        except Exception as e:
            raise UserError(f"Error decoding file: {str(e)}")
        
        # Read Excel file
        try:
            # Try openpyxl first (for .xlsx)
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(file_content), data_only=True)
            sheet = workbook.active
            rows = list(sheet.values)
        except Exception:
            try:
                # Fallback to xlrd (for .xls)
                workbook = xlrd.open_workbook(file_contents=file_content)
                sheet = workbook.sheet_by_index(0)
                rows = [sheet.row_values(i) for i in range(sheet.nrows)]
            except Exception as e:
                raise UserError(f"Error reading Excel file: {str(e)}\n"
                              f"Please ensure the file is a valid Excel file (.xls or .xlsx)")
        
        if not rows or len(rows) < 2:
            raise UserError("The Excel file is empty or has no data rows.")
        
        # Parse data
        headers = [str(h).strip() if h else '' for h in rows[0]]
        data_rows = rows[1:]
        
        # Process import
        results = self._process_import(headers, data_rows)
        
        # Show results
        return self._show_import_results(results)

    def _process_import(self, headers, data_rows):
        """Process the import data"""
        results = {
            'created': 0,
            'updated': 0,
            'skipped': 0,
            'errors': []
        }
        
        ProductPackaging = self.env['product.packaging']
        ProductProduct = self.env['product.product']
        PackageType = self.env['stock.package.type']
        StockRoute = self.env['stock.route']
        
        # Map headers to field names
        header_map = {
            'product code/id': 'product_ref',
            'packaging name': 'name',
            'sequence': 'sequence',
            'contained quantity': 'qty',
            'barcode': 'barcode',
            'package type': 'package_type',
            'sales': 'sales',
            'purchase': 'purchase',
            'routes': 'routes',
        }
        
        # Find column indices
        col_indices = {}
        for idx, header in enumerate(headers):
            header_lower = header.lower().strip().rstrip('*')
            if header_lower in header_map:
                col_indices[header_map[header_lower]] = idx
        
        # Validate required columns
        required_fields = ['product_ref', 'name', 'qty']
        missing_fields = [f for f in required_fields if f not in col_indices]
        if missing_fields:
            raise UserError(f"Missing required columns: {', '.join(missing_fields)}")
        
        # Track products for replace_existing
        products_to_clear = set()
        
        # Build allowed product IDs for filtering
        allowed_product_ids = None
        if self.filter_product_id:
            allowed_product_ids = {self.filter_product_id.id}
        elif self.filter_product_tmpl_id:
            allowed_product_ids = set(self.filter_product_tmpl_id.product_variant_ids.ids)
        
        for row_num, row in enumerate(data_rows, start=2):
            try:
                # Skip empty rows
                if not any(row):
                    continue
                
                # Extract values
                vals = {}
                for field, idx in col_indices.items():
                    if idx < len(row):
                        vals[field] = row[idx]
                
                # Find product
                product_ref = str(vals.get('product_ref', '')).strip()
                if not product_ref:
                    results['errors'].append(f"Row {row_num}: Product Code/ID is required")
                    results['skipped'] += 1
                    continue
                
                # Check if it's an ID (in brackets) or default_code
                if product_ref.startswith('[') and product_ref.endswith(']'):
                    try:
                        product_id = int(product_ref[1:-1])
                        product = ProductProduct.browse(product_id)
                    except ValueError:
                        results['errors'].append(f"Row {row_num}: Invalid product ID format: {product_ref}")
                        results['skipped'] += 1
                        continue
                else:
                    product = ProductProduct.search([('default_code', '=', product_ref)], limit=1)
                
                if not product.exists():
                    results['errors'].append(f"Row {row_num}: Product not found: {product_ref}")
                    results['skipped'] += 1
                    continue
                
                # Check if product is in allowed list (when filtering)
                if allowed_product_ids is not None and product.id not in allowed_product_ids:
                    results['errors'].append(
                        f"Row {row_num}: Product '{product.display_name}' is not in the filtered product list - skipped"
                    )
                    results['skipped'] += 1
                    continue
                
                # Clear existing packaging if needed (only once per product)
                if self.replace_existing and product.id not in products_to_clear:
                    product.packaging_ids.unlink()
                    products_to_clear.add(product.id)
                
                # Prepare packaging values
                packaging_vals = self._prepare_packaging_values(vals, product, row_num, results)
                
                if not packaging_vals:
                    results['skipped'] += 1
                    continue
                
                # Find package type
                if vals.get('package_type'):
                    package_type_name = str(vals['package_type']).strip()
                    package_type = PackageType.search([('name', '=', package_type_name)], limit=1)
                    if package_type:
                        packaging_vals['package_type_id'] = package_type.id
                    else:
                        results['errors'].append(f"Row {row_num}: Package type not found: {package_type_name}")
                
                # Find routes
                if vals.get('routes'):
                    routes_str = str(vals['routes']).strip()
                    if routes_str:
                        route_names = [r.strip() for r in routes_str.split(',') if r.strip()]
                        route_ids = []
                        for route_name in route_names:
                            route = StockRoute.search([
                                ('name', '=', route_name),
                                ('packaging_selectable', '=', True)
                            ], limit=1)
                            if route:
                                route_ids.append(route.id)
                            else:
                                results['errors'].append(
                                    f"Row {row_num}: Route not found or not selectable: {route_name}"
                                )
                        if route_ids:
                            packaging_vals['route_ids'] = [(6, 0, route_ids)]
                
                # Create or update packaging
                if self.import_type in ['update', 'both']:
                    # Try to find existing packaging by name and product
                    existing = ProductPackaging.search([
                        ('product_id', '=', product.id),
                        ('name', '=', packaging_vals['name'])
                    ], limit=1)
                    
                    if existing:
                        existing.write(packaging_vals)
                        results['updated'] += 1
                        continue
                
                if self.import_type in ['create', 'both']:
                    ProductPackaging.create(packaging_vals)
                    results['created'] += 1
                else:
                    results['skipped'] += 1
                    results['errors'].append(
                        f"Row {row_num}: Packaging '{packaging_vals['name']}' not found for update"
                    )
                
            except Exception as e:
                results['errors'].append(f"Row {row_num}: {str(e)}")
                results['skipped'] += 1
        
        return results

    def _prepare_packaging_values(self, vals, product, row_num, results):
        """Prepare values for packaging creation/update"""
        packaging_vals = {
            'product_id': product.id,
            'company_id': self.company_id.id,
        }
        
        # Name (required)
        name = str(vals.get('name', '')).strip()
        if not name:
            results['errors'].append(f"Row {row_num}: Packaging name is required")
            return False
        packaging_vals['name'] = name
        
        # Quantity (required)
        try:
            qty = float(vals.get('qty', 0))
            if qty <= 0:
                results['errors'].append(f"Row {row_num}: Contained quantity must be greater than 0")
                return False
            packaging_vals['qty'] = qty
        except (ValueError, TypeError):
            results['errors'].append(f"Row {row_num}: Invalid contained quantity: {vals.get('qty')}")
            return False
        
        # Sequence (optional)
        if vals.get('sequence'):
            try:
                packaging_vals['sequence'] = int(vals['sequence'])
            except (ValueError, TypeError):
                results['errors'].append(f"Row {row_num}: Invalid sequence value: {vals.get('sequence')}")
        
        # Barcode (optional)
        if vals.get('barcode'):
            barcode = str(vals['barcode']).strip()
            if barcode:
                packaging_vals['barcode'] = barcode
        
        # Sales (optional, default True)
        packaging_vals['sales'] = self._parse_boolean(vals.get('sales'), True)
        
        # Purchase (optional, default True)
        packaging_vals['purchase'] = self._parse_boolean(vals.get('purchase'), True)
        
        return packaging_vals

    def _parse_boolean(self, value, default=False):
        """Parse boolean value from Excel"""
        if value is None or value == '':
            return default
        
        value_str = str(value).strip().upper()
        
        if value_str in ['TRUE', 'YES', '1', 'T', 'Y']:
            return True
        elif value_str in ['FALSE', 'NO', '0', 'F', 'N']:
            return False
        
        return default

    def _show_import_results(self, results):
        """Show import results to user"""
        message_lines = []
        
        if results['created']:
            message_lines.append(f"✓ Created: {results['created']} packaging record(s)")
        
        if results['updated']:
            message_lines.append(f"✓ Updated: {results['updated']} packaging record(s)")
        
        if results['skipped']:
            message_lines.append(f"⚠ Skipped: {results['skipped']} row(s)")
        
        if results['errors']:
            message_lines.append(f"\n⚠ Errors ({len(results['errors'])}):")
            # Show first 10 errors
            for error in results['errors'][:10]:
                message_lines.append(f"  • {error}")
            if len(results['errors']) > 10:
                message_lines.append(f"  ... and {len(results['errors']) - 10} more errors")
        
        message = '\n'.join(message_lines)
        
        # Determine notification type
        if results['errors'] and not (results['created'] or results['updated']):
            notif_type = 'danger'
            title = 'Import Failed'
        elif results['errors']:
            notif_type = 'warning'
            title = 'Import Completed with Errors'
        else:
            notif_type = 'success'
            title = 'Import Successful'
        
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': title,
                'message': message,
                'type': notif_type,
                'sticky': True if results['errors'] else False,
            }
        }
